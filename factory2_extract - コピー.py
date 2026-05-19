"""
第2工場 生産機械抽出ツール
ダブルクリックで起動します。Python + openpyxl のみ必要。
"""

import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from collections import Counter
from datetime import datetime

# ─── openpyxl チェック ───────────────────────────────────────────
try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess, sys as _sys
    root = tk.Tk(); root.withdraw()
    if messagebox.askyesno(
        "ライブラリ不足",
        "openpyxl がインストールされていません。\n今すぐインストールしますか？"
    ):
        subprocess.check_call([_sys.executable, "-m", "pip", "install", "openpyxl"])
        messagebox.showinfo("完了", "インストール完了。ツールを再起動してください。")
    sys.exit(0)


# ─── 抽出ロジック ────────────────────────────────────────────────
TARGET_GROUP = "第2工場"

COL = {
    "表示項目":   2,
    "OrderNo":    3,
    "機種グループ": 9,
    "生産月":    10,
    "機種":      11,
    "機種名":    12,
    "号機":      13,
    "納入先":    14,
    "立会":      15,
    "仕様":      16,
    "完成要求日": 19,
}

COLOR_HEADER_BG = "305496"
COLOR_SUBTOTAL  = "D9E1F2"


def _thin_border():
    s = Side(border_style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_style(cell):
    cell.font = Font(name="Arial", bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", start_color=COLOR_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _thin_border()


def _data_style(cell, align="left"):
    cell.font = Font(name="Arial", size=10)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = _thin_border()


def _find_data_start(ws):
    for r in range(1, 200):
        for c in range(1, 30):
            v = ws.cell(row=r, column=c).value
            if v and str(v).strip() == "OrderNo":
                return r + 1
    raise ValueError("ヘッダー行（OrderNo）が見つかりません。")


def _extract_records(ws, data_start):
    seen = set()
    records = []
    for r in range(data_start, ws.max_row + 1):
        if ws.cell(row=r, column=COL["機種グループ"]).value != TARGET_GROUP:
            continue
        disp = ws.cell(row=r, column=COL["表示項目"]).value
        if disp in (99, 2):
            continue
        order = ws.cell(row=r, column=COL["OrderNo"]).value
        if order in seen:
            continue
        seen.add(order)
        records.append({
            "OrderNo":    order,
            "生産月":     ws.cell(row=r, column=COL["生産月"]).value,
            "機種":       ws.cell(row=r, column=COL["機種"]).value,
            "機種名":     ws.cell(row=r, column=COL["機種名"]).value,
            "号機":       ws.cell(row=r, column=COL["号機"]).value,
            "納入先":     ws.cell(row=r, column=COL["納入先"]).value,
            "立会":       ws.cell(row=r, column=COL["立会"]).value,
            "仕様":       ws.cell(row=r, column=COL["仕様"]).value,
            "完成要求日": ws.cell(row=r, column=COL["完成要求日"]).value,
        })
    records.sort(key=lambda x: (x["生産月"] or 0, x["機種"] or "", x["号機"] or 0))
    return records


def _write_summary(ws, records):
    ws.title = "集計"
    ws["A1"] = f"{TARGET_GROUP} 生産機械一覧（集計）"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws.merge_cells("A1:C1")

    ws["A3"] = "■ 機種別台数"
    ws["A3"].font = Font(name="Arial", bold=True, size=11)

    for col, label in zip("ABC", ["機種", "機種名（内訳）", "台数"]):
        _header_style(ws[f"{col}4"])
        ws[f"{col}4"].value = label

    machine_names = {}
    for rec in records:
        machine_names.setdefault(rec["機種"], []).append(rec["機種名"])

    type_counter = Counter(r["機種"] for r in records)
    r = 5
    for machine, count in sorted(type_counter.items(), key=lambda x: -x[1]):
        names_str = "、".join(
            f"{n}({c}台)" for n, c in Counter(machine_names[machine]).most_common()
        )
        for c, v in [(1, machine), (2, names_str), (3, count)]:
            _data_style(ws.cell(row=r, column=c, value=v),
                        align="right" if c == 3 else "left")
        r += 1

    for c, v in [(1, "合計"), (3, f"=SUM(C5:C{r-1})")]:
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", start_color=COLOR_SUBTOTAL)
        cell.border = _thin_border()
        cell.alignment = Alignment(horizontal="right" if c == 3 else "left")
    ws.cell(row=r, column=2).fill = PatternFill("solid", start_color=COLOR_SUBTOTAL)
    ws.cell(row=r, column=2).border = _thin_border()

    mr = r + 3
    ws.cell(row=mr, column=1, value="■ 生産月別台数").font = Font(name="Arial", bold=True, size=11)
    hr = mr + 1
    for col, label in zip("AB", ["生産月", "台数"]):
        _header_style(ws[f"{col}{hr}"])
        ws[f"{col}{hr}"].value = label

    month_counter = Counter(rec["生産月"] for rec in records)
    dr = hr + 1
    for month, count in sorted(month_counter.items()):
        _data_style(ws.cell(row=dr, column=1, value=month))
        _data_style(ws.cell(row=dr, column=2, value=count), align="right")
        dr += 1
    for c, v in [(1, "合計"), (2, f"=SUM(B{hr+1}:B{dr-1})")]:
        cell = ws.cell(row=dr, column=c, value=v)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", start_color=COLOR_SUBTOTAL)
        cell.border = _thin_border()
        cell.alignment = Alignment(horizontal="right" if c == 2 else "left")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 10


def _write_detail(ws, records):
    ws.title = "明細"
    headers = [
        ("No.", 6), ("OrderNo", 14), ("生産月", 12), ("機種", 10),
        ("機種名", 22), ("号機", 8), ("納入先", 50),
        ("立会", 8), ("仕様", 35), ("完成要求日", 14),
    ]
    for col, (label, width) in enumerate(headers, 1):
        _header_style(ws.cell(row=1, column=col, value=label))
        ws.column_dimensions[get_column_letter(col)].width = width

    for ri, rec in enumerate(records, start=2):
        for col, val in enumerate([
            ri - 1, rec["OrderNo"], rec["生産月"], rec["機種"],
            rec["機種名"], rec["号機"], rec["納入先"],
            rec["立会"], rec["仕様"], rec["完成要求日"],
        ], 1):
            cell = ws.cell(row=ri, column=col, value=val)
            _data_style(cell, align="right" if col in (1, 6) else "left")
        ws.cell(row=ri, column=10).number_format = "yyyy/mm/dd"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def run_extraction(input_path, output_path, log):
    """抽出処理本体（別スレッドで呼ばれる）"""
    wb_src = openpyxl.load_workbook(input_path, data_only=True, keep_vba=False)
    if "main" not in wb_src.sheetnames:
        raise ValueError("'main' シートが見つかりません。")

    ws = wb_src["main"]
    log("ヘッダー行を検索中...")
    data_start = _find_data_start(ws)
    log(f"データ開始行: {data_start} 行目")

    log(f"'{TARGET_GROUP}' のデータを抽出中...")
    records = _extract_records(ws, data_start)
    log(f"抽出件数: {len(records)} 台")

    if not records:
        raise ValueError(f"'{TARGET_GROUP}' のデータが見つかりませんでした。")

    log("Excelファイルを作成中...")
    wb = Workbook()
    _write_summary(wb.active, records)
    _write_detail(wb.create_sheet(), records)
    wb.save(output_path)

    # 集計結果を返す
    return records


# ─── GUI ─────────────────────────────────────────────────────────
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("第2工場 生産機械抽出ツール")
        self.resizable(False, False)
        self._build_ui()
        self._center()

    def _center(self):
        self.update_idletasks()
        w, h = 560, 440
        x = (self.winfo_screenwidth()  - w) // 2
        y = (self.winfo_screenheight() - h) // 2
        self.geometry(f"{w}x{h}+{x}+{y}")

    def _build_ui(self):
        PAD = {"padx": 16, "pady": 6}

        # ── タイトルバー ──
        title_frame = tk.Frame(self, bg="#305496")
        title_frame.pack(fill="x")
        tk.Label(
            title_frame, text="第2工場 生産機械抽出ツール",
            bg="#305496", fg="white",
            font=("Arial", 13, "bold"), pady=10
        ).pack()

        # ── 入力ファイル ──
        f1 = tk.LabelFrame(self, text="入力ファイル（.xlsm / .xlsx）",
                           font=("Arial", 9), padx=8, pady=6)
        f1.pack(fill="x", **PAD, pady=(14, 4))

        self.input_var = tk.StringVar()
        tk.Entry(f1, textvariable=self.input_var, width=52,
                 font=("Arial", 9)).grid(row=0, column=0, padx=(0, 6))
        tk.Button(f1, text="参照...", command=self._browse_input,
                  width=8).grid(row=0, column=1)

        # ── 出力ファイル ──
        f2 = tk.LabelFrame(self, text="出力ファイル（.xlsx）",
                           font=("Arial", 9), padx=8, pady=6)
        f2.pack(fill="x", **PAD, pady=4)

        self.output_var = tk.StringVar()
        tk.Entry(f2, textvariable=self.output_var, width=52,
                 font=("Arial", 9)).grid(row=0, column=0, padx=(0, 6))
        tk.Button(f2, text="参照...", command=self._browse_output,
                  width=8).grid(row=0, column=1)

        # ── 実行ボタン ──
        self.run_btn = tk.Button(
            self, text="▶  抽 出 実 行",
            font=("Arial", 12, "bold"),
            bg="#217346", fg="white",
            activebackground="#19593a", activeforeground="white",
            relief="flat", pady=8,
            command=self._on_run
        )
        self.run_btn.pack(fill="x", padx=16, pady=(8, 4))

        # ── プログレスバー ──
        self.progress = ttk.Progressbar(self, mode="indeterminate")
        self.progress.pack(fill="x", padx=16, pady=(0, 4))

        # ── ログ ──
        log_frame = tk.LabelFrame(self, text="実行ログ",
                                  font=("Arial", 9), padx=6, pady=4)
        log_frame.pack(fill="both", expand=True, padx=16, pady=(0, 12))

        self.log_text = tk.Text(
            log_frame, height=8, font=("Consolas", 9),
            state="disabled", bg="#f8f8f8", relief="flat"
        )
        sb = tk.Scrollbar(log_frame, command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=sb.set)
        self.log_text.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

    # ── ファイル選択 ──
    def _browse_input(self):
        path = filedialog.askopenfilename(
            title="入力ファイルを選択",
            filetypes=[("Excel ファイル", "*.xlsm *.xlsx"), ("すべて", "*.*")]
        )
        if not path:
            return
        self.input_var.set(path)
        # 出力パスを自動設定
        p = Path(path)
        today = datetime.now().strftime("%Y%m%d")
        auto_out = p.parent / f"{p.stem}_第2工場抽出_{today}.xlsx"
        self.output_var.set(str(auto_out))

    def _browse_output(self):
        path = filedialog.asksaveasfilename(
            title="出力先を選択",
            defaultextension=".xlsx",
            filetypes=[("Excel ファイル", "*.xlsx")]
        )
        if path:
            self.output_var.set(path)

    # ── ログ書き込み ──
    def _log(self, msg: str):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    # ── 実行 ──
    def _on_run(self):
        input_path  = self.input_var.get().strip()
        output_path = self.output_var.get().strip()

        if not input_path:
            messagebox.showwarning("入力エラー", "入力ファイルを選択してください。")
            return
        if not Path(input_path).exists():
            messagebox.showerror("エラー", f"ファイルが見つかりません:\n{input_path}")
            return
        if not output_path:
            messagebox.showwarning("入力エラー", "出力ファイルのパスを指定してください。")
            return

        self.run_btn.configure(state="disabled")
        self.progress.start(10)
        self._log("処理を開始します...")

        def worker():
            try:
                records = run_extraction(
                    Path(input_path), Path(output_path), self._log
                )
                self.after(0, lambda: self._on_success(records, output_path))
            except Exception as e:
                self.after(0, lambda: self._on_error(str(e)))

        threading.Thread(target=worker, daemon=True).start()

    def _on_success(self, records, output_path):
        self.progress.stop()
        self.run_btn.configure(state="normal")

        type_counter = Counter(r["機種"] for r in records)
        summary = "\n".join(f"  {m}: {c}台" for m, c in
                            sorted(type_counter.items(), key=lambda x: -x[1]))
        self._log(f"━━ 完了 ━━  合計 {len(records)} 台")
        self._log(f"出力先: {output_path}")

        if messagebox.askyesno(
            "完了",
            f"抽出完了しました。\n\n合計: {len(records)} 台\n\n{summary}\n\n"
            "出力ファイルを開きますか？"
        ):
            import subprocess, os
            subprocess.Popen(["start", "", output_path], shell=True)

    def _on_error(self, msg):
        self.progress.stop()
        self.run_btn.configure(state="normal")
        self._log(f"[ERROR] {msg}")
        messagebox.showerror("エラー", msg)


if __name__ == "__main__":
    App().mainloop()
