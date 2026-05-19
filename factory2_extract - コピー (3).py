import openpyxl
from collections import Counter
from datetime import datetime

TARGET_GROUP = "第2工場"

COL = {
    "OrderNo": 3,
    "機種グループ": 9,
    "機種": 11,
}

# -----------------------------
# ログ表示
# -----------------------------
def log(msg):
    t = datetime.now().strftime("%H:%M:%S")
    print(f"[{t}] {msg}")

# -----------------------------
# データ開始位置探す
# -----------------------------
def find_start(ws):
    log("① ヘッダー検索")
    for r in range(1, 200):
        for c in range(1, 30):
            val = ws.cell(r, c).value
            if val and str(val).strip() == "OrderNo":
                return r + 1
    raise ValueError("OrderNoが見つからない")

# -----------------------------
# メイン処理
# -----------------------------
def run(input_file, output_file):

    log("=== 処理開始 ===")
    log(f"入力: {input_file}")
    log(f"出力: {output_file}")

    wb = openpyxl.load_workbook(input_file)
    ws = wb.active

    start_row = find_start(ws)

    counter = Counter()
    total = 0

    log("② 抽出開始")

    for row in range(start_row, ws.max_row + 1):
        group = ws.cell(row, COL["機種グループ"]).value
        machine = ws.cell(row, COL["機種"]).value

        if group == TARGET_GROUP and machine:
            counter[machine] += 1
            total += 1

    # -----------------------------
    # コンソール出力（全文）
    # -----------------------------
    log("③ 集計結果")

    print("\n=== 第2工場 機種別台数 ===")
    for machine, count in sorted(counter.items(), key=lambda x: x[1], reverse=True):
        print(f"{machine}: {count}台")

    print(f"\n合計: {total}台")

    # -----------------------------
    # Excel出力
    # -----------------------------
    log("④ Excel出力")

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = "集計結果"

    # ヘッダー
    out_ws["A1"] = "機種"
    out_ws["B1"] = "台数"

    row = 2
    for machine, count in sorted(counter.items(), key=lambda x: x[1], reverse=True):
        out_ws.cell(row, 1, machine)
        out_ws.cell(row, 2, count)
        row += 1

    # 合計行
    out_ws.cell(row, 1, "合計")
    out_ws.cell(row, 2, total)

    out_wb.save(output_file)

    log("✅ 完了")


# -----------------------------
# 実行
# -----------------------------
if __name__ == "__main__":
    input_file = "input.xlsm"
    output_file = "result.xlsx"

    run(input_file, output_file)
